"""Events summary (published vs draft), division summaries, and the Excel export."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from scheduling_engine import Event, League, fixtures
from scheduling_engine.export import _sheet_name, schedule_workbook
from scheduling_engine.matchups import Matchup
from scheduling_engine.report import fairness_report
from scheduling_engine.solver import SolveOptions, solve
from scheduling_engine.summary import events_summary


@pytest.fixture(scope="module")
def season() -> tuple[League, list[Event]]:
    lg = fixtures.small_8()
    return lg, solve(lg, options=SolveOptions(time_limit_s=10, workers=4)).events


def test_summary_counts_published_and_draft(season: tuple[League, list[Event]]) -> None:
    lg, events = season
    moved = events[0].model_copy(update={"start": events[0].start.replace(hour=14)})
    now = [moved, *events[1:]]
    rows = {r.team_id: r for r in events_summary(lg, now, published=events)}
    for t in lg.teams:
        r = rows[t.id]
        assert r.games == r.published_games + r.unpublished_games == 7
        assert r.home + r.away == r.games
        changed = t.id in moved.team_ids
        assert r.unpublished_games == (1 if changed else 0)


def test_summary_counts_games_waiting_for_a_spot(season: tuple[League, list[Event]]) -> None:
    lg, events = season
    e = events[0]
    assert e.away_team_id is not None
    waiting = [Matchup(e.id, e.division_id, e.home_team_id, e.away_team_id, 0)]
    rows = {r.team_id: r for r in events_summary(lg, events[1:], waiting=waiting)}
    assert rows[e.home_team_id].waiting == 1 and rows[e.home_team_id].games == 6
    assert all(r.unpublished_games == r.games for r in rows.values()), "nothing published yet"


def test_division_summary_and_what_stands_out(season: tuple[League, list[Event]]) -> None:
    lg, events = season
    rep = fairness_report(lg, events)
    (div,) = rep.divisions
    assert div.teams == 8 and div.games == 28 and div.fewest_games == div.most_games == 7
    assert div.verdict.startswith("Every team plays 7 games, home and away within 1")
    assert rep.highlights and all(h.endswith(".") for h in rep.highlights)

    short = fairness_report(lg, events[3:])
    assert "play fewer games than the target." in " ".join(short.highlights)
    assert "(target 7)" in short.divisions[0].verdict
    assert all(f.kind for f in short.flags)


def test_workbook_has_schedule_summary_and_a_sheet_per_team(
    season: tuple[League, list[Event]],
) -> None:
    lg, events = season
    data = schedule_workbook(lg, events, title="Spring")
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames[:2] == ["Schedule", "Summary"]
    assert len(wb.sheetnames) == 2 + len(lg.teams)
    ws = wb["Schedule"]
    assert ws["A1"].value == "Spring"
    header = [c.value for c in ws[3]]
    assert header[:7] == ["Date", "Day", "Start", "End", "Division", "Home", "Away"]
    assert ws.max_row == 3 + len(events)
    team = lg.teams[0]
    assert wb[team.name].max_row == 3 + sum(1 for e in events if team.id in e.team_ids)


def test_sheet_names_are_valid_and_unique() -> None:
    used: set[str] = set()
    a = _sheet_name("Hawks: U11/Blue [A] with a very long team name", used)
    b = _sheet_name("Hawks: U11/Blue [A] with a very long team name", used)
    assert len(a) <= 31 and len(b) <= 31 and a != b
    assert not any(ch in a for ch in "[]:*?/\\")
