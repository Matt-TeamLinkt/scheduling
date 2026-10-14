"""Read a CSV or Excel file into rows of cells, whatever shape it arrives in."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time

Cell = str | int | float | date | datetime | time | None


@dataclass
class RawRow:
    number: int  # the row number the scheduler sees in their spreadsheet (1-based)
    cells: list[Cell]


class UnreadableFile(Exception):
    """Raised with a plain-language reason the file can't be read at all."""


def read_rows(data: bytes, filename: str) -> list[RawRow]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    if name.endswith(".xls"):
        raise UnreadableFile(
            "This is an older Excel file (.xls). Open it in Excel or Google Sheets and "
            "save it as .xlsx or CSV, then try again."
        )
    if data[:4] == b"PK\x03\x04":  # an .xlsx with the wrong extension
        return _read_xlsx(data)
    return _read_csv(data)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_csv(data: bytes) -> list[RawRow]:
    text = _decode(data)
    if not text.strip():
        raise UnreadableFile("The file is empty.")
    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[RawRow] = []
    for i, cells in enumerate(csv.reader(io.StringIO(text), dialect), 1):
        if any(c.strip() for c in cells):
            rows.append(RawRow(number=i, cells=[c.strip() or None for c in cells]))
    return rows


def _read_xlsx(data: bytes) -> list[RawRow]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise UnreadableFile("This Excel file couldn't be opened. Is it damaged?") from exc
    # First sheet that has anything in it.
    for ws in wb.worksheets:
        rows: list[RawRow] = []
        for i, values in enumerate(ws.iter_rows(values_only=True), 1):
            cells: list[Cell] = [_xlsx_cell(v) for v in values]
            if any(c is not None for c in cells):
                rows.append(RawRow(number=i, cells=cells))
        if rows:
            return rows
    raise UnreadableFile("The spreadsheet is empty.")


def _xlsx_cell(v: object) -> Cell:
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip() or None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float, date, datetime, time)):
        return v
    return str(v)
