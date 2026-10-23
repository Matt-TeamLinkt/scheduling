"""Spreadsheet export: the whole schedule, a summary, and one sheet per team.

    data = schedule_workbook(league, events, waiting, title="Winter 2027")
    open("schedule.xlsx", "wb").write(data)

Dates and times are real Excel dates and times, so the scheduler can sort, filter and
print. Column names are plain, and the "Schedule" sheet uses the same columns the
importer will read back.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .index import LeagueIndex
from .matchups import Matchup
from .model import Event, League
from .summary import events_summary

Value = str | int | float | date | time | None
HEADER_FILL = PatternFill("solid", fgColor="E8EEFF")
BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def schedule_workbook(
    league: League,
    events: Sequence[Event],
    waiting: Sequence[Matchup] = (),
    published: Sequence[Event] = (),
    title: str | None = None,
) -> bytes:
    ix = LeagueIndex(league)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Schedule"
    ordered = sorted(events, key=lambda e: (e.start, ix.place(e.slot_id), e.id))

    _table(
        ws,
        ["Date", "Day", "Start", "End", "Division", "Home", "Away", "Venue", "Field", "Locked"],
        [
            [
                e.start.date(),
                f"{e.start:%a}",
                e.start.time(),
                e.end.time(),
                _division(ix, e.division_id),
                ix.team_name(e.home_team_id),
                ix.team_name(e.away_team_id) if e.away_team_id else "Practice",
                *_venue_field(ix, e.slot_id),
                "Yes" if e.locked else "",
            ]
            for e in ordered
        ],
        title,
    )

    rows = events_summary(league, events, published, waiting)
    _table(
        wb.create_sheet("Summary"),
        [
            "Team",
            "Division",
            "Games",
            "Target",
            "Home",
            "Away",
            "Published games",
            "Unpublished games",
            "Practices",
            "Needs a spot",
        ],
        [
            [
                r.name,
                r.division,
                r.games,
                r.target,
                r.home,
                r.away,
                r.published_games,
                r.unpublished_games,
                r.practices,
                r.waiting,
            ]
            for r in sorted(rows, key=lambda r: (r.division, r.name))
        ],
    )

    if waiting:
        _table(
            wb.create_sheet("Needs a spot"),
            ["Division", "Home", "Away"],
            [
                [
                    _division(ix, m.division_id),
                    ix.team_name(m.home_team_id),
                    ix.team_name(m.away_team_id),
                ]
                for m in waiting
            ],
        )

    used = {s.lower() for s in wb.sheetnames}
    for team in sorted(league.teams, key=lambda t: (_division(ix, t.division_id), t.name)):
        mine = [e for e in ordered if team.id in e.team_ids]
        name = _sheet_name(team.name, used)
        _table(
            wb.create_sheet(name),
            ["Date", "Day", "Start", "End", "Opponent", "Home or away", "Venue", "Field"],
            [
                [
                    e.start.date(),
                    f"{e.start:%a}",
                    e.start.time(),
                    e.end.time(),
                    _opponent(ix, e, team.id),
                    ""
                    if e.away_team_id is None
                    else "Home"
                    if e.home_team_id == team.id
                    else "Away",
                    *_venue_field(ix, e.slot_id),
                ]
                for e in mine
            ],
            f"{team.name} · {_division(ix, team.division_id)}",
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _division(ix: LeagueIndex, division_id: str) -> str:
    d = ix.divisions.get(division_id)
    return d.name if d else ""


def _venue_field(ix: LeagueIndex, slot_id: str) -> list[str]:
    slot = ix.slots.get(slot_id)
    surface = ix.surfaces.get(slot.surface_id) if slot else None
    venue = ix.venues.get(surface.venue_id) if surface else None
    return [venue.name if venue else "", surface.name if surface else ""]


def _opponent(ix: LeagueIndex, e: Event, team_id: str) -> str:
    if e.away_team_id is None:
        return "Practice"
    return ix.team_name(e.away_team_id if e.home_team_id == team_id else e.home_team_id)


def _sheet_name(name: str, used: set[str]) -> str:
    """Excel sheet names: at most 31 characters, no []:*?/\\, unique (ignoring case)."""
    base = BAD_SHEET_CHARS.sub("", name).strip()[:31] or "Team"
    out, n = base, 2
    while out.lower() in used:
        suffix = f" ({n})"
        out = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(out.lower())
    return out


def _table(
    ws: Worksheet, header: list[str], rows: list[list[Value]], title: str | None = None
) -> None:
    top = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        top = 3
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    widths = [len(h) for h in header]
    for r, row in enumerate(rows, start=top + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, datetime):
                cell.number_format = "yyyy-mm-dd h:mm AM/PM"
            elif isinstance(v, date):
                cell.number_format = "ddd mmm d, yyyy"
                widths[c - 1] = max(widths[c - 1], 16)
            elif isinstance(v, time):
                cell.number_format = "h:mm AM/PM"
                widths[c - 1] = max(widths[c - 1], 9)
            else:
                widths[c - 1] = max(widths[c - 1], len(str(v)))
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = min(40, w + 2)
    ws.freeze_panes = f"A{top + 1}"
    if rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(header))}{top + len(rows)}"
